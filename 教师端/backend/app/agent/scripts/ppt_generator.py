# -*- coding: utf-8 -*-
r"""多版式 PPT 生成器（孔孟智能体 · 备课 Agent 生成脚本）
=========================================================
设计目标（质量硬要求）：
  1) 内容充实详尽：每个内容页必须包含具体知识点讲解、案例说明与总结提炼，
     由 validate_slides() 硬校验（要点数量/字数/总结字段），不达标拒绝生成；
  2) 版式多样化：标题页 / 目录页 / 概念页（定义框+要点） / 案例页（图文混排：
     左要点右函数图像） / 图表页（学情数据大图+结论） / 分栏对比页 / 表格页 /
     总结页（要点卡片） / 结束页 —— 九种版式，单一版式占比超 50% 直接拒绝；
  3) 排版规范：竹青-秋香赭-暖墨配色、渐变噪声背景、LaTeX 公式渲染为图片、
     Markdown 符号一律清理、页码与页脚齐全。

页面 Schema（gen_ppt 工具的 slidesJson 每一项）：
  {
    "layout":  "title|toc|concept|case|chart|twocolumn|table|summary|end",
    "title":   "页面标题",
    "subtitle":"副标题（title/end 页）",
    "points":  ["要点1（尽量具体，含知识点讲解）", ...],
    "box":     "concept 页左侧定义/概念框内容",
    "left":    {"heading": "...", "points": [...]},   # twocolumn 左栏
    "right":   {"heading": "...", "points": [...]},   # twocolumn 右栏
    "caseText":"case 页案例题干/情境（≥40 字）",
    "plotSpec":{"kind":"function|bar|line", "expr":"x**3-3*x", "range":[-2,2],
                "labels":[...], "values":[...], "ylabel":"..."},  # 图文混排/图表页
    "chartData":{"labels":[...], "values":[...], "title":"..."},      # chart 页
    "table":   {"headers":[...], "rows":[[...]]},                     # table 页
    "summary": "本页总结提炼（≥15 字）",
    "note":    "讲稿备注",
  }

CLI：python ppt_generator.py <slides.json> <out.pptx>
"""
from __future__ import annotations

import io
import json
import os
import re
import sys
from typing import Any, Dict, List, Optional, Tuple

_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

import workflow  # noqa: E402  复用 _latex_img/_mixed_line_img/_strip_md_inline/_mpl/_safe_filename

LAYOUTS = ("title", "toc", "concept", "case", "chart", "twocolumn", "table", "summary", "end")

# ---------------------------------------------------------------------------
# 质量硬校验
# ---------------------------------------------------------------------------
def validate_slides(slides: List[Dict[str, Any]]) -> Dict[str, Any]:
    issues: List[str] = []
    if not isinstance(slides, list) or len(slides) < 12:
        issues.append(f"页数不足：{len(slides) if isinstance(slides, list) else 0} < 12")
        return {"ok": False, "issues": issues}
    # 版式合法
    for i, s in enumerate(slides, 1):
        if s.get("layout") not in LAYOUTS:
            issues.append(f"第{i}页 layout 非法：{s.get('layout')}")
    content = [s for s in slides if s.get("layout") not in ("title", "toc", "end")]
    # 1) 内容充实度：每个内容页要点数/字数
    for i, s in enumerate(slides, 1):
        lay = s.get("layout")
        if lay in ("title", "toc", "end"):
            continue
        pts = s.get("points") or []
        if lay == "twocolumn":
            for side in ("left", "right"):
                p2 = (s.get(side) or {}).get("points") or []
                chars = sum(len(str(p)) for p in p2)
                if len(p2) < 2 or chars < 50:
                    issues.append(f"第{i}页 {side} 栏要点不足（{len(p2)} 条/{chars} 字，需≥2 条且≥50 字）")
        elif lay == "table":
            t = s.get("table") or {}
            if not t.get("headers") or not t.get("rows") or len(t["rows"]) < 3:
                issues.append(f"第{i}页 表格数据不足（需表头+≥3 行）")
        else:
            chars = sum(len(str(p)) for p in pts)
            if len(pts) < 2 or chars < 60:
                issues.append(f"第{i}页 要点单薄（{len(pts)} 条/{chars} 字，需≥2 条且≥60 字）")
        if lay == "case" and len(str(s.get("caseText") or "")) < 40:
            issues.append(f"第{i}页 案例题干过短（≥40 字）")
        if lay in ("concept", "case", "summary") and len(str(s.get("summary") or "")) < 15:
            issues.append(f"第{i}页 缺少总结提炼（summary ≥15 字）")
    # 2) 版式多样性：单一版式占比 ≤50%，图表页/分栏或表格页至少各 1
    from collections import Counter
    cnt = Counter(s.get("layout") for s in content)
    if content:
        top_lay, top_n = cnt.most_common(1)[0]
        if top_n / len(content) > 0.5:
            issues.append(f"版式单一：{top_lay} 占内容页 {top_n}/{len(content)}（>50%）")
    distinct = len(cnt)
    if len(slides) >= 12 and distinct < 6:
        issues.append(f"版式种类不足：{distinct} 种（≥12 页时需 ≥6 种）")
    if not cnt.get("chart"):
        issues.append("缺少图表页（学情数据可视化）")
    if not (cnt.get("twocolumn") or cnt.get("table")):
        issues.append("缺少分栏对比页或表格页")
    if not cnt.get("case"):
        issues.append("缺少案例页（图文混排）")
    # 3) 空洞词检测
    for i, s in enumerate(slides, 1):
        blob = json.dumps(s, ensure_ascii=False)
        for w in ("此处讲解", "略。", "等等。", "（占位）"):
            if w in blob:
                issues.append(f"第{i}页 含占位性表述「{w}」")
    return {"ok": not issues, "issues": issues, "layoutStats": dict(cnt),
            "contentPages": len(content), "total": len(slides)}


# ---------------------------------------------------------------------------
# 渲染基础设施
# ---------------------------------------------------------------------------
JADE = (79, 122, 102)     # 竹青主色
TAN = (168, 129, 78)      # 秋香赭点缀
INK = (47, 42, 34)        # 暖墨
LIGHT = (247, 246, 241)   # 米白
PANEL = (233, 238, 233)   # 浅青面板


def _bg_image(dpi: int = 110, tint: Tuple[int, int, int] = (226, 233, 228)):
    """全页渐变+噪声背景（保证质感与体积）。"""
    import numpy as np
    plt = workflow._mpl()
    w, h = 13.333, 7.5
    fig = plt.figure(figsize=(w, h), dpi=dpi)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.axis("off")
    n_y, n_x = 300, 540
    yy, xx = np.mgrid[0:n_y, 0:n_x]
    base = np.array(LIGHT, dtype=float)
    tn = np.array(tint, dtype=float)
    mix = (xx / n_x * 0.55 + yy / n_y * 0.45)[..., None]
    img = base * (1 - mix) + tn * mix
    rng = np.random.default_rng(42)
    img = np.clip(img + rng.normal(0, 1.6, img.shape), 0, 255).astype("uint8")
    ax.imshow(img, aspect="auto")
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi)
    plt.close(fig)
    buf.seek(0)
    return buf


def _chart_image(spec: Dict[str, Any], dpi: int = 160):
    """根据 chartData/plotSpec 生成图表图片（bar/line/function）。"""
    import numpy as np
    plt = workflow._mpl()
    kind = str(spec.get("kind") or "bar")
    fig, ax = plt.subplots(figsize=(7.2, 4.6), dpi=dpi)
    fig.patch.set_facecolor("#F7F6F1")
    ax.set_facecolor("#F7F6F1")
    if kind == "function":
        expr = str(spec.get("expr") or "x")
        rng = spec.get("range") or [-5, 5]
        env = {"x": None, "np": np, "sqrt": np.sqrt, "abs": np.abs, "sin": np.sin,
               "cos": np.cos, "tan": np.tan, "exp": np.exp, "log": np.log, "pi": np.pi}
        xs = np.linspace(float(rng[0]), float(rng[1]), 400)
        env["x"] = xs
        ys = eval(expr, {"__builtins__": {}}, env)  # 受限求值：仅 numpy 函数与 x
        ax.plot(xs, ys, color="#4F7A66", lw=2.6)
        ax.axhline(0, color="#8A8578", lw=0.8)
        ax.axvline(0, color="#8A8578", lw=0.8)
        ax.grid(alpha=0.3)
        ax.set_title(str(spec.get("title") or f"y = {expr}"), color="#2F2A22")
    else:
        labels = [str(x) for x in (spec.get("labels") or [])]
        values = [float(v) for v in (spec.get("values") or [])]
        colors = ["#B4504A" if v < 60 else ("#E2A33C" if v < 85 else "#4F7A66") for v in values]
        if kind == "line":
            ax.plot(labels, values, marker="o", color="#4F7A66", lw=2.2)
        else:
            bars = ax.bar(labels, values, color=colors, edgecolor="none")
            for b, v in zip(bars, values):
                ax.text(b.get_x() + b.get_width() / 2, v + 1.5, f"{v:g}", ha="center",
                        fontsize=10, color="#2F2A22")
        ax.set_ylim(0, 110)
        ax.axhline(60, color="#B4504A", lw=1.0, ls="--", alpha=0.7)
        ax.axhline(85, color="#4F7A66", lw=1.0, ls="--", alpha=0.7)
        ax.set_ylabel(str(spec.get("ylabel") or "掌握率(%)"), color="#2F2A22")
        ax.set_title(str(spec.get("title") or ""), color="#2F2A22")
        ax.grid(axis="y", alpha=0.3)
        for lb in ax.get_xticklabels():
            lb.set_color("#2F2A22")
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight", facecolor=fig.get_facecolor())
    plt.close(fig)
    buf.seek(0)
    return buf


def _render_pptx(slides: List[Dict[str, Any]], path: str, meta: Dict[str, Any]) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    topic = str(meta.get("topic") or "")
    subject = str(meta.get("subject") or "")
    grade = str(meta.get("grade") or "")

    def _rgb(c):
        return RGBColor(*c)

    def _new_slide(tint=(226, 233, 228)):
        s = prs.slides.add_slide(blank)
        try:
            s.shapes.add_picture(_bg_image(tint=tint), 0, 0, width=prs.slide_width, height=prs.slide_height)
        except Exception:
            s.background.fill.solid()
            s.background.fill.fore_color.rgb = _rgb(LIGHT)
        return s

    def _txt(slide, l, t, w, h, lines, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, line_spacing=1.15):
        tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        if isinstance(lines, str):
            lines = [lines]
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.space_after = Pt(6)
            r = p.add_run()
            r.text = workflow._strip_md_inline(str(ln))
            f = r.font
            f.size = Pt(size)
            f.color.rgb = _rgb(color)
            f.bold = bold
        return tb

    def _bullets(slide, l, t, w, h, points, size=16, color=INK, mark="• "):
        rows = []
        for p in points or []:
            txt = workflow._strip_md_inline(str(p))
            if workflow._has_latex(txt):
                # 含中文的公式行用混排渲染（_latex_img 不含 CJK 字形会出方块）
                rows.append(("IMG_MIX" if re.search(r"[\u4e00-\u9fff]", txt) else "IMG", txt))
            else:
                rows.append(("TXT", txt))
        tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        idx = 0
        for kind, txt in rows:
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.space_after = Pt(7)
            r = p.add_run()
            if kind == "TXT":
                r.text = mark + txt
                r.font.size = Pt(size)
                r.font.color.rgb = _rgb(color)
            else:
                r.text = "· "
                r.font.size = Pt(size)
                r.font.color.rgb = _rgb(color)
                try:
                    img = workflow._mixed_line_img(txt) if kind == "IMG_MIX" else workflow._latex_img(txt)
                    slide.shapes.add_picture(img, Inches(l + 0.28), Inches(t + idx * 0.52 + 0.02), height=Inches(0.44))
                except Exception:
                    r.text = mark + txt
                    r.font.size = Pt(size - 2)
            idx += 1
        return tb

    def _panel(slide, l, t, w, h, fill=PANEL, line=None):
        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = _rgb(fill)
        if line:
            sh.line.color.rgb = _rgb(line)
            sh.line.width = Pt(1.2)
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    def _title_bar(slide, title, sub=None, color=JADE, size=28):
        _txt(slide, 0.6, 0.3, 11.0, 0.9, title, size=size, color=color, bold=True)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(1.18), Inches(1.6), Emu(40000))
        bar.fill.solid()
        bar.fill.fore_color.rgb = _rgb(TAN)
        bar.line.fill.background()
        if sub:
            _txt(slide, 10.2, 0.42, 2.6, 0.6, sub, size=12, color=(138, 133, 120), align=PP_ALIGN.RIGHT)
        _page_footer(slide)

    def _page_footer(slide):
        n = len(prs.slides._sldIdLst)
        _txt(slide, 12.35, 7.05, 0.7, 0.35, f"{n:02d}", size=11, color=(138, 133, 120), align=PP_ALIGN.RIGHT)
        _txt(slide, 0.62, 7.05, 5.0, 0.35, f"孔孟教研智能体 · {subject}{(' ' + grade) if grade else ''}",
             size=10, color=(160, 156, 144))

    def _summary_strip(slide, text, t=6.35):
        _panel(slide, 0.6, t, 12.13, 0.72, fill=(236, 232, 219))
        _txt(slide, 0.85, t + 0.08, 11.7, 0.6, [("【总结提炼】" + workflow._strip_md_inline(str(text)))],
             size=14, color=INK, bold=True)

    # ---------------- 各版式 ----------------
    for s in slides:
        lay = s.get("layout", "concept")
        title = str(s.get("title") or "")
        if lay == "title":
            sl = _new_slide(tint=(219, 229, 222))
            _panel(sl, 4.4, 2.0, 4.53, 0.09, fill=TAN)
            _txt(sl, 1.2, 2.5, 10.93, 1.4, title, size=44, color=JADE, bold=True, align=PP_ALIGN.CENTER)
            _txt(sl, 1.2, 4.0, 10.93, 0.7, str(s.get("subtitle") or f"高中{subject} · {grade}"),
                 size=20, color=INK, align=PP_ALIGN.CENTER)
            _txt(sl, 1.2, 6.3, 10.93, 0.5, "孔孟教研智能体 · 基于学情数据总线自动生成",
                 size=13, color=(138, 133, 120), align=PP_ALIGN.CENTER)
            continue
        if lay == "toc":
            sl = _new_slide()
            _title_bar(sl, "目录", "CONTENTS")
            items = s.get("points") or []
            col_x = [0.9, 6.9]
            per = (len(items) + 1) // 2
            for i, it in enumerate(items):
                cx = col_x[i // per]
                cy = 1.7 + (i % per) * 0.92
                circ = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx), Inches(cy), Inches(0.52), Inches(0.52))
                circ.fill.solid()
                circ.fill.fore_color.rgb = _rgb(JADE)
                circ.line.fill.background()
                tf = circ.text_frame
                tf.paragraphs[0].alignment = PP_ALIGN.CENTER
                r = tf.paragraphs[0].add_run()
                r.text = str(i + 1)
                r.font.size = Pt(16)
                r.font.bold = True
                r.font.color.rgb = _rgb((255, 255, 255))
                _txt(sl, cx + 0.72, cy + 0.03, 5.4, 0.6, it, size=17, color=INK, bold=True)
            continue
        if lay == "concept":
            sl = _new_slide()
            _title_bar(sl, title)
            box = str(s.get("box") or "")
            _panel(sl, 0.6, 1.55, 4.3, 4.6, fill=PANEL, line=JADE)
            _txt(sl, 0.85, 1.7, 3.9, 0.5, "概念要点", size=15, color=JADE, bold=True)
            _txt(sl, 0.85, 2.2, 3.9, 3.8, [ln for ln in box.split("\n") if ln.strip()],
                 size=14, color=INK)
            _bullets(sl, 5.25, 1.6, 7.45, 4.55, s.get("points") or [], size=15)
            if s.get("summary"):
                _summary_strip(sl, s["summary"])
            continue
        if lay == "case":
            sl = _new_slide()
            _title_bar(sl, title, "图文混排")
            _panel(sl, 0.6, 1.5, 6.3, 4.7, fill=(240, 239, 231))
            _txt(sl, 0.8, 1.6, 5.95, 1.15, "【案例】" + workflow._strip_md_inline(str(s.get("caseText") or "")),
                 size=13.5, color=INK)
            _bullets(sl, 0.8, 2.85, 5.95, 3.2, s.get("points") or [], size=13.5)
            plot = s.get("plotSpec") or {}
            if plot:
                try:
                    sl.shapes.add_picture(_chart_image(plot), Inches(7.15), Inches(1.5), width=Inches(5.55))
                except Exception:
                    pass
            if s.get("summary"):
                _summary_strip(sl, s["summary"])
            continue
        if lay == "chart":
            sl = _new_slide()
            _title_bar(sl, title, "数据看板")
            cd = dict(s.get("chartData") or {})
            cd.setdefault("kind", "bar")
            try:
                sl.shapes.add_picture(_chart_image(cd), Inches(0.6), Inches(1.5), width=Inches(8.1))
            except Exception:
                pass
            _panel(sl, 8.95, 1.5, 3.78, 4.7, fill=PANEL, line=JADE)
            _txt(sl, 9.15, 1.65, 3.4, 0.5, "数据结论", size=15, color=JADE, bold=True)
            _bullets(sl, 9.15, 2.2, 3.4, 3.9, s.get("points") or [], size=13)
            if s.get("summary"):
                _summary_strip(sl, s["summary"])
            continue
        if lay == "twocolumn":
            sl = _new_slide()
            _title_bar(sl, title, "对比辨析")
            for side, (px, hd_color, fill) in {
                "left": (0.6, JADE, (232, 238, 232)),
                "right": (6.75, TAN, (242, 233, 221)),
            }.items():
                data = s.get(side) or {}
                _panel(sl, px, 1.55, 5.95, 4.65, fill=fill)
                _txt(sl, px + 0.25, 1.7, 5.4, 0.6, str(data.get("heading") or ("左栏" if side == "left" else "右栏")),
                     size=17, color=hd_color, bold=True)
                _bullets(sl, px + 0.25, 2.35, 5.45, 3.7, data.get("points") or [], size=13.5)
            if s.get("summary"):
                _summary_strip(sl, s["summary"])
            continue
        if lay == "table":
            sl = _new_slide()
            _title_bar(sl, title, "结构化呈现")
            t = s.get("table") or {}
            headers = [workflow._strip_md_inline(str(x)) for x in (t.get("headers") or [])]
            rows = t.get("rows") or []
            n_rows = len(rows) + 1
            gt = sl.shapes.add_table(n_rows, len(headers), Inches(0.6), Inches(1.6),
                                     Inches(12.13), Inches(min(0.62 * n_rows, 4.6))).table
            for c, htxt in enumerate(headers):
                cell = gt.cell(0, c)
                cell.text = htxt
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(JADE)
                pr = cell.text_frame.paragraphs[0]
                pr.alignment = PP_ALIGN.CENTER
                pr.runs[0].font.size = Pt(15)
                pr.runs[0].font.bold = True
                pr.runs[0].font.color.rgb = _rgb((255, 255, 255))
            for ri, row in enumerate(rows, 1):
                for ci, val in enumerate(row):
                    cell = gt.cell(ri, ci)
                    cell.text = workflow._strip_md_inline(str(val))
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = _rgb(LIGHT if ri % 2 else (238, 240, 234))
                    for p in cell.text_frame.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(13)
                            r.font.color.rgb = _rgb(INK)
            if s.get("summary"):
                _summary_strip(sl, s["summary"], t=6.35)
            continue
        if lay == "summary":
            sl = _new_slide()
            _title_bar(sl, title, "本课提炼")
            pts = list(s.get("points") or [])[:4]
            card_h = min(1.05, 4.4 / max(len(pts), 1))
            for i, p in enumerate(pts):
                cy = 1.6 + i * (card_h + 0.18)
                _panel(sl, 0.9, cy, 11.5, card_h, fill=(238, 236, 226) if i % 2 else PANEL)
                num = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.1), Inches(cy + card_h / 2 - 0.21),
                                          Inches(0.42), Inches(0.42))
                num.fill.solid()
                num.fill.fore_color.rgb = _rgb(TAN)
                num.line.fill.background()
                tf = num.text_frame
                tf.paragraphs[0].alignment = PP_ALIGN.CENTER
                r = tf.paragraphs[0].add_run()
                r.text = str(i + 1)
                r.font.size = Pt(14)
                r.font.bold = True
                r.font.color.rgb = _rgb((255, 255, 255))
                _txt(sl, 1.75, cy + 0.08, 10.4, card_h, p, size=15, color=INK, bold=True)
            if s.get("summary"):
                _summary_strip(sl, s["summary"])
            continue
        if lay == "end":
            sl = _new_slide(tint=(219, 229, 222))
            _txt(sl, 1.2, 2.3, 10.93, 1.1, "本课到此结束 · 感谢聆听", size=36, color=JADE, bold=True,
                 align=PP_ALIGN.CENTER)
            pts = s.get("points") or []
            if pts:
                _panel(sl, 3.4, 3.9, 6.53, 1.7, fill=(236, 232, 219))
                _txt(sl, 3.65, 4.0, 6.0, 0.5, "课后思考", size=15, color=TAN, bold=True)
                _txt(sl, 3.65, 4.5, 6.1, 1.0, pts[:2], size=13, color=INK)
            continue
        # 未知版式兜底 → concept
        sl = _new_slide()
        _title_bar(sl, title)
        _bullets(sl, 0.8, 1.6, 11.7, 4.6, s.get("points") or [], size=15)

    prs.save(path)


def generate_ppt(slides: List[Dict[str, Any]], out_path: str, meta: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    """生成入口：先校验后渲染。返回质量统计；不达标抛 ValueError。"""
    meta = meta or {}
    quality = validate_slides(slides)
    if not quality["ok"]:
        raise ValueError("PPT 质量校验未通过：" + "；".join(quality["issues"][:6]))
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    _render_pptx(slides, out_path, meta)
    size = os.path.getsize(out_path)
    if size < 500_000:  # 体积下限：背景/图片确保图文质感
        raise ValueError(f"PPT 体积过小（{size} 字节 < 500KB），版面可能空洞")
    quality["path"] = out_path
    quality["size"] = size
    return quality


def main() -> None:  # CLI：独立运行生成
    if len(sys.argv) < 3:
        print("用法：python ppt_generator.py <slides.json> <out.pptx> [meta.json]")
        sys.exit(1)
    with open(sys.argv[1], "r", encoding="utf-8") as f:
        slides = json.load(f)
    meta = {}
    if len(sys.argv) > 3 and os.path.exists(sys.argv[3]):
        with open(sys.argv[3], "r", encoding="utf-8") as f:
            meta = json.load(f)
    q = generate_ppt(slides, sys.argv[2], meta)
    print(json.dumps(q, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
